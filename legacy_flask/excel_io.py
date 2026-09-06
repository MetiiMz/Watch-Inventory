# -*- coding: utf-8 -*-
"""ورودی/خروجی CSV و Excel — خروجی گرفتن و ورود داده با openpyxl."""

import csv
import io
import os

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from database import get_db
from jalali import parse_jalali_date, gregorian_to_jalali

PRODUCT_HEADERS = [
    "نام ساعت", "رفرنس", "کد انبار سایت", "کد دفتر فروشگاه",
    "برند", "قیمت خرید", "قیمت فروش", "وضعیت",
    "تأمین‌کننده", "تاریخ خرید", "یادداشت",
]
PRODUCT_FIELDS = [
    "name", "reference", "website_code", "office_code", "brand",
    "purchase_price", "sale_price", "available", "supplier",
    "purchase_date", "notes",
]

REPAIR_HEADERS = [
    "نام ساعت", "کد ساعت", "ایراد", "تاریخ تحویل", "تاریخ بازگشت",
    "نام مشتری", "شماره تماس", "زیر گارانتی", "وضعیت", "هزینه تعمیر", "یادداشت",
]

TRACKING_HEADERS = [
    "نام ساعت یا قطعه", "کد / مشخصات", "نام مشتری", "شماره تماس",
    "قیمت", "وضعیت", "یادداشت",
]
TRACKING_FIELDS = [
    "item_name", "item_code", "customer_name", "customer_phone",
    "price", "status", "notes",
]

PAYMENTS_HEADERS = [
    "نام محصول", "نام مشتری", "شماره تماس", "مبلغ کل",
    "پرداخت‌شده", "مانده", "تاریخ", "یادداشت",
]
PAYMENTS_FIELDS = [
    "product_name", "customer_name", "customer_phone",
    "total_amount", "paid_amount", "pay_date", "notes",
]

STATUS_FA = {
    "received": "دریافت شده",
    "in_progress": "در حال تعمیر",
    "waiting_parts": "انتظار قطعه",
    "done": "آماده تحویل",
    "delivered": "تحویل شده",
}
STATUS_EN = {v: k for k, v in STATUS_FA.items()}

TRACKING_STATUS_FA = {
    "new": "جدید",
    "ordered": "سفارش داده شد",
    "found": "پیدا شد / آماده",
    "delivered": "تحویل شد",
    "cancelled": "لغو شد",
}
TRACKING_STATUS_EN = {v: k for k, v in TRACKING_STATUS_FA.items()}

PURCHASE_TYPE_FA = {"person": "حضوری", "online": "آنلاین"}
PAYMENT_TYPE_FA = {"cash": "نقدی", "card": "کارت به کارت", "deposit": "بیعانه"}


def _jdate_from_iso(iso):
    if not iso:
        return ""
    try:
        y, m, d = (int(x) for x in str(iso)[:10].split("-"))
        jy, jm, jd = gregorian_to_jalali(y, m, d)
        return f"{jy}/{jm:02d}/{jd:02d}"
    except (ValueError, AttributeError):
        return ""


def _money(value):
    if value is None:
        return 0
    try:
        f = float(value)
    except (TypeError, ValueError):
        return 0
    return int(f) if f == int(f) else round(f, 2)


# ---------------------------------------------------------------- products

def _product_rows():
    rows = []
    with get_db() as db:
        recs = db.execute("SELECT * FROM products ORDER BY office_code COLLATE NOCASE").fetchall()
    for r in recs:
        rows.append([
            r["name"], r["reference"], r["website_code"], r["office_code"],
            r["brand"], _money(r["purchase_price"]), _money(r["sale_price"]),
            "موجود" if r["available"] else "ناموجود",
            r["supplier"], _jdate_from_iso(r["purchase_date"]),
            r["notes"],
        ])
    return rows


def _export_csv(path, headers, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(headers)
        w.writerows(rows)


def _export_xlsx(path, headers, rows, sheet_title="گزارش"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.sheet_view.rightToLeft = True
    _write_table(ws, headers, rows)
    wb.save(path)


def export_products(path, fmt):
    if fmt == "xlsx":
        _export_xlsx(path, PRODUCT_HEADERS, _product_rows(), "انبار")
    else:
        _export_csv(path, PRODUCT_HEADERS, _product_rows())


# ---------------------------------------------------------------- repairs

def _repair_rows():
    rows = []
    with get_db() as db:
        recs = db.execute("SELECT * FROM repairs ORDER BY id DESC").fetchall()
    for r in recs:
        rows.append([
            r["watch_name"], r["watch_code"], r["issue"],
            _jdate_from_iso(r["delivery_date"]), _jdate_from_iso(r["return_date"]),
            r["customer_name"], r["customer_phone"],
            "بله" if r["is_warranty"] else "خیر",
            STATUS_FA.get(r["status"], r["status"]),
            _money(r["repair_price"]), r["notes"],
        ])
    return rows


def export_repairs(path, fmt):
    if fmt == "xlsx":
        _export_xlsx(path, REPAIR_HEADERS, _repair_rows(), "تعمیرات")
    else:
        _export_csv(path, REPAIR_HEADERS, _repair_rows())


# ---------------------------------------------------------------- tracking

def _tracking_rows():
    rows = []
    with get_db() as db:
        recs = db.execute("SELECT * FROM tracking ORDER BY id DESC").fetchall()
    for r in recs:
        rows.append([
            r["item_name"], r["item_code"], r["customer_name"],
            r["customer_phone"], _money(r["price"]),
            TRACKING_STATUS_FA.get(r["status"], r["status"]), r["notes"],
        ])
    return rows


def export_tracking(path, fmt):
    if fmt == "xlsx":
        _export_xlsx(path, TRACKING_HEADERS, _tracking_rows(), "پیگیری")
    else:
        _export_csv(path, TRACKING_HEADERS, _tracking_rows())


# ---------------------------------------------------------------- payments

def _payment_rows():
    rows = []
    with get_db() as db:
        recs = db.execute("SELECT * FROM payments ORDER BY id DESC").fetchall()
    for r in recs:
        rows.append([
            r["product_name"], r["customer_name"], r["customer_phone"],
            _money(r["total_amount"]), _money(r["paid_amount"]),
            _money(r["total_amount"] - r["paid_amount"]),
            _jdate_from_iso(r["pay_date"]), r["notes"],
        ])
    return rows


def export_payments(path, fmt):
    if fmt == "xlsx":
        _export_xlsx(path, PAYMENTS_HEADERS, _payment_rows(), "پرداخت‌ها")
    else:
        _export_csv(path, PAYMENTS_HEADERS, _payment_rows())


# ---------------------------------------------------------------- sales

SALES_HEADERS = ["کد دفتر فروشگاه", "کد انبار سایت", "رفرنس", "نام ساعت", "برند",
                 "قیمت خرید", "قیمت فروش", "قیمت نهایی", "سود",
                 "تاریخ فروش", "خریدار", "شماره تماس", "نوع فروش",
                 "نقدی", "کارت‌خوان", "کارت به کارت", "نوع پرداخت", "یادداشت"]


def export_sales(path, fmt):
    rows = []
    with get_db() as db:
        recs = db.execute(
            """
            SELECT s.*, p.name AS product_name, p.website_code, p.office_code,
                   p.reference, p.brand, p.purchase_price AS p_purchase_price
            FROM sales s LEFT JOIN products p ON p.id = s.product_id
            ORDER BY s.sale_date DESC, s.id DESC
            """
        ).fetchall()
    for r in recs:
        final = r["final_price"] or r["sale_price"] or 0
        rows.append([
            r["office_code"] or "", r["website_code"] or "", r["reference"] or "",
            r["product_name"] or "", r["brand"] or "",
            _money(r["p_purchase_price"] if "p_purchase_price" in r.keys() else r["purchase_price"]),
            _money(r["sale_price"]), _money(final),
            _money(r["profit"]),
            _jdate_from_iso(r["sale_date"]), r["customer"], r["customer_phone"] or "",
            PURCHASE_TYPE_FA.get(r["sale_type"], "حضوری"),
            _money(r["paid_cash"]), _money(r["paid_pos"]), _money(r["paid_card2card"]),
            PAYMENT_TYPE_FA.get(r["payment_type"], "نقدی"),
            r["notes"],
        ])
    if fmt == "xlsx":
        _export_xlsx(path, SALES_HEADERS, rows, "فروش‌ها")
    else:
        _export_csv(path, SALES_HEADERS, rows)


# ---------------------------------------------------------------- table writer

def _write_table(ws, headers, rows):
    header_fill = PatternFill("solid", fgColor="1F4FD8")
    header_font = Font(bold=True, color="FFFFFF", name="Vazirmatn", size=11)
    body_font = Font(name="Vazirmatn", size=10)
    align = Alignment(horizontal="center", vertical="center")

    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align
    ws.row_dimensions[1].height = 26

    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = body_font
            cell.alignment = align
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 18


# ---------------------------------------------------------------- import

def import_products(filepath, update_existing=True):
    """
    ورود محصولات از CSV یا XLSX.
    بازگشت: (ok, آمار, خطاها[])
    """
    rows, err = _read_table(filepath, PRODUCT_HEADERS)
    if err:
        return False, None, err

    stats = {"added": 0, "updated": 0, "skipped": 0}
    errors = []
    with get_db() as db:
        for i, row in enumerate(rows, 2):  # شروع از سطر ۲ (بعد از سرستون)
            try:
                rec = dict(zip(PRODUCT_FIELDS, row))
                name = str(rec.get("name") or "").strip()
                if not name:
                    stats["skipped"] += 1
                    continue

                office_code = str(rec.get("office_code") or "").strip()
                website_code = str(rec.get("website_code") or "").strip()
                if not office_code and not website_code:
                    errors.append(f"سطر {i}: «{name}» بدون کد انبار سایت و کد دفتر است؛ رد شد")
                    stats["skipped"] += 1
                    continue
                if not office_code:
                    office_code = website_code
                if not website_code:
                    website_code = office_code

                try:
                    purchase_price = float(str(rec.get("purchase_price") or 0).replace(",", "") or 0)
                except ValueError:
                    purchase_price = 0
                try:
                    sale_price = float(str(rec.get("sale_price") or 0).replace(",", "") or 0)
                except ValueError:
                    sale_price = 0

                # وضعیت در فایل: «موجود» / «ناموجود» (یا تعداد در فایل‌های قدیمی)
                raw_status = str(rec.get("available") or rec.get("quantity") or "").strip()
                if raw_status in ("موجود", "1", 1, "true", "True"):
                    available = 1
                elif raw_status in ("ناموجود", "0", 0, "false", "False"):
                    available = 0
                else:
                    try:
                        available = 1 if int(float(raw_status or 1)) > 0 else 0
                    except ValueError:
                        available = 1

                purchase_date = parse_jalali_date(rec.get("purchase_date")) or ""

                values = {
                    "name": name,
                    "reference": str(rec.get("reference") or "").strip(),
                    "office_code": office_code,
                    "website_code": website_code,
                    "brand": str(rec.get("brand") or "").strip(),
                    "purchase_price": purchase_price,
                    "sale_price": sale_price,
                    "available": available,
                    "supplier": str(rec.get("supplier") or "").strip(),
                    "purchase_date": purchase_date,
                    "notes": str(rec.get("notes") or "").strip(),
                }

                existing = db.execute(
                    "SELECT id FROM products WHERE office_code = ? OR website_code = ?",
                    (office_code, website_code),
                ).fetchone()

                if existing:
                    if update_existing:
                        sets = ", ".join(f"{k} = ?" for k in values)
                        db.execute(
                            f"UPDATE products SET {sets} WHERE id = ?",
                            (*values.values(), existing["id"]),
                        )
                        stats["updated"] += 1
                    else:
                        stats["skipped"] += 1
                else:
                    cols = ", ".join(values.keys())
                    marks = ", ".join("?" for _ in values)
                    db.execute(
                        f"INSERT INTO products({cols}) VALUES({marks})",
                        tuple(values.values()),
                    )
                    stats["added"] += 1
            except Exception as e:  # noqa: BLE001 — هر سطر مستقل است
                errors.append(f"سطر {i}: {e}")
                stats["skipped"] += 1

    return True, stats, errors


def _read_table(filepath, expected_headers):
    """CSV یا XLSX را می‌خواند و لیستی از ردیف‌ها (بدون سرستون) برمی‌گرداند."""
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".csv":
        return _read_csv(filepath, expected_headers)
    if ext in (".xlsx", ".xlsm", ".xls"):
        return _read_xlsx(filepath, expected_headers)
    return None, "فرمت فایل پشتیبانی نمی‌شود (فقط csv و xlsx)"


def _read_csv(filepath, expected_headers):
    raw = open(filepath, "rb").read()
    if raw.startswith(b"\xef\xbb\xbf"):
        raw = raw[3:]
    text = None
    for enc in ("utf-8-sig", "utf-8", "cp1256"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        return None, "انکودینگ فایل قابل خواندن نبود"
    sample = text[:4096]
    delim = ";" if sample.count(";") > sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delim)
    all_rows = [r for r in reader if any((c or "").strip() for c in r)]
    if not all_rows:
        return None, "فایل خالی است"
    header = [h.strip() for h in all_rows[0]]
    if not _headers_match(header, expected_headers):
        return None, "سرستون‌های فایل با قالب استاندارد مطابقت ندارد. لطفاً ابتدا خروجی نمونه بگیرید."
    return all_rows[1:], None


def _read_xlsx(filepath, expected_headers):
    wb = load_workbook(filepath, data_only=True)
    ws = wb.active
    all_rows = []
    for row in ws.iter_rows(values_only=True):
        if row is None:
            continue
        vals = [("" if v is None else v) for v in row]
        if any(str(v).strip() for v in vals):
            all_rows.append(vals)
    if not all_rows:
        return None, "فایل خالی است"
    header = [str(h).strip() for h in all_rows[0]]
    if not _headers_match(header, expected_headers):
        return None, "سرستون‌های فایل با قالب استاندارد مطابقت ندارد. لطفاً ابتدا خروجی نمونه بگیرید."
    return all_rows[1:], None


def _headers_match(header, expected):
    h = [str(x).strip() for x in header if str(x).strip()]
    e = [str(x).strip() for x in expected]
    common = set(h) & set(e)
    return len(common) >= max(3, len(e) - 3)
